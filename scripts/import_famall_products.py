from __future__ import annotations

import json
import re
import shutil
import unicodedata
from dataclasses import dataclass
from difflib import SequenceMatcher
from pathlib import Path
from urllib.parse import urljoin
from urllib.request import Request, urlopen

from docx import Document
from docx.oxml.ns import qn
from lxml import html as lxml_html
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
PRICE_PDF = Path("/Users/viktoria/Desktop/ПРАЙС-ЛИСТ FAMALL- новый.pdf")
CATALOG_DOCX_CANDIDATES = [
    Path("/Users/viktoria/Desktop/FAMALL_каталог_без_цен.docx"),
    Path("/Users/viktoria/Downloads/Telegram Desktop/FAMALL_каталог_без_цен.docx"),
    Path("/Users/viktoria/Desktop/FAMALL_полный_каталог.docx"),
    Path("/Users/viktoria/Downloads/Telegram Desktop/FAMALL_полный_каталог.docx"),
    Path("/Users/viktoria/Desktop/FAMALL_каталог_с_реальными_изображениями.docx"),
]
SITE_URL = "https://famallnetwork62.orgs.biz/#products"
SITE_BASE = "https://famallnetwork62.orgs.biz/"
TAPLINK_URL = "https://famall-ufa.taplink.ws/"
TAPLINK_BASE = "https://famall-ufa.taplink.ws/"
TAPLINK_IMAGE_BASE = "https://taplink.st/p/"
PRODUCT_IMAGES_DIR = ROOT / "assets/images/products"
CATALOG_IMAGES_DIR = PRODUCT_IMAGES_DIR / "catalog"
PLACEHOLDER_IMAGES_DIR = PRODUCT_IMAGES_DIR / "placeholders"
PRODUCTS_JSON = ROOT / "data/products.json"
CATEGORIES_JSON = ROOT / "data/categories.json"
UNMATCHED_JSON = ROOT / "data/import-unmatched.json"
REVIEW_JSON = ROOT / "data/import-review.json"
REPORT_JSON = ROOT / "data/import-report.json"
MISSING_IMAGES_JSON = ROOT / "data/missing-images-report.json"
IMPORT_CACHE = ROOT / "data/import-cache"

BRANDS = [
    "BERCLEAN",
    "YIJIAN",
    "LIMANCY",
    "PREDAWN",
    "MIAOROU",
    "MIUBABY",
    "OKFAD",
    "BOCARE",
    "SUTING",
    "MOREFUTURE",
    "CARICH",
    "ILIFE",
    "BEIJING",
    "ХАН",
]


@dataclass
class PriceItem:
    index: int
    name: str
    brand: str
    category: str
    volume: str
    partner_price: int
    retail_price: int
    pv: float


def fetch(url: str) -> bytes:
    request = Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urlopen(request, timeout=30) as response:
        return response.read()


def cached_fetch(url: str, name: str) -> bytes:
    IMPORT_CACHE.mkdir(parents=True, exist_ok=True)
    cache_file = IMPORT_CACHE / name
    if cache_file.exists():
        return cache_file.read_bytes()
    data = fetch(url)
    cache_file.write_bytes(data)
    return data


def normalize_text(value: str) -> str:
    value = unicodedata.normalize("NFKC", value).replace("ё", "е").replace("Ё", "Е")
    value = re.sub(r"[^\wа-яА-Я]+", " ", value.lower(), flags=re.UNICODE)
    return re.sub(r"\s+", " ", value).strip()


def slugify(value: str) -> str:
    translit = {
        "а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "е": "e", "ё": "e", "ж": "zh", "з": "z",
        "и": "i", "й": "y", "к": "k", "л": "l", "м": "m", "н": "n", "о": "o", "п": "p", "р": "r",
        "с": "s", "т": "t", "у": "u", "ф": "f", "х": "h", "ц": "c", "ч": "ch", "ш": "sh", "щ": "sch",
        "ъ": "", "ы": "y", "ь": "", "э": "e", "ю": "yu", "я": "ya",
    }
    text = "".join(translit.get(char, char) for char in value.lower())
    text = re.sub(r"[^a-z0-9]+", "-", text)
    return text.strip("-")[:90]


def xml_escape(value: str) -> str:
    return (
        value.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


def shorten(value: str, limit: int) -> str:
    value = re.sub(r"\s+", " ", value).strip()
    return value if len(value) <= limit else f"{value[: limit - 1].rstrip()}…"


def wrap_text(value: str, line_limit: int = 26, max_lines: int = 4) -> list[str]:
    words = re.sub(r"\s+", " ", value).strip().split()
    lines: list[str] = []
    current = ""
    for word in words:
        next_line = f"{current} {word}".strip()
        if len(next_line) <= line_limit:
            current = next_line
            continue
        if current:
            lines.append(current)
        current = word
        if len(lines) == max_lines - 1:
            break
    if current and len(lines) < max_lines:
        lines.append(current)
    if not lines:
        lines = [shorten(value, line_limit)]
    return [shorten(line, line_limit) for line in lines[:max_lines]]


def detect_brand(name: str) -> str:
    normalized = normalize_text(name).upper()
    for brand in BRANDS:
        if normalize_text(brand).upper() in normalized:
            return brand.title() if brand not in {"OKFAD", "BOCARE", "SUTING"} else brand
    return "FAMALL"


def split_volume(name: str) -> tuple[str, str]:
    patterns = [
        r"(\d+(?:[.,]\d+)?\s*(?:мл|мг|гр|г|g|шт|штук|рулон(?:а|ов)?|предмет(?:ов|а)?|пакетик(?:ов|а)?|лист(?:ов)?|мм|кг)\.?(?:\s*[*xх/]\s*\d+\s*(?:шт|штук|пакетик(?:ов)?|лист(?:ов)?)?)?)$",
        r"(р-р\s*[\d*хx/ ]+(?:,\s*)?\d*\s*(?:штук|шт)?)$",
        r"(\d+\s*[-–]\s*\d+\s*кг)$",
        r"(\d+)$",
    ]
    cleaned = re.sub(r"\s+", " ", name).strip()
    for pattern in patterns:
        match = re.search(pattern, cleaned, flags=re.IGNORECASE)
        if match:
            volume = match.group(1).strip(" .,")
            product_name = cleaned[: match.start()].strip(" .,")
            if len(product_name) >= 4:
                return product_name, volume
    return cleaned, ""


def parse_price_pdf() -> list[PriceItem]:
    text = "\n".join((page.extract_text() or "") for page in PdfReader(PRICE_PDF).pages)
    lines = [line.strip() for line in text.splitlines() if line.strip()]

    items: list[PriceItem] = []
    category = ""
    current_number = 1
    buffer: list[str] = []

    def flush() -> None:
        nonlocal buffer, current_number
        if not buffer:
            return
        raw = " ".join(buffer)
        raw = re.sub(r"\s+", " ", raw).strip()
        raw = re.sub(rf"^{current_number}\s*", "", raw)
        match = re.search(r"(.+?)\s+(\d+)\s+(\d+)\s+(\d+(?:,\d+)?)$", raw)
        if not match:
            buffer = []
            current_number += 1
            return
        name_with_volume = match.group(1).strip()
        name, volume = split_volume(name_with_volume)
        items.append(
            PriceItem(
                index=current_number,
                name=name,
                brand=detect_brand(name),
                category=category or "FAMALL",
                volume=volume,
                partner_price=int(match.group(2)),
                retail_price=int(match.group(3)),
                pv=float(match.group(4).replace(",", ".")),
            )
        )
        buffer = []
        current_number += 1

    for line in lines:
        if line.startswith("ПРАЙС") or line.startswith("Наименование") or line.startswith("Цена "):
            continue

        category_line = not re.match(r"^\d{1,3}\s*", line) and not re.search(r"\d+\s+\d+\s+\d+(?:,\d+)?$", line)
        if category_line and line.upper() == line:
            flush()
            category = line.replace("СЕРИЯ", "").replace('"', "").strip(" :")
            continue

        number_match = re.match(r"^(\d{1,3})(.*)$", line)
        if number_match:
            number = int(number_match.group(1))
            if number == current_number:
                flush()
                buffer = [line]
                if re.search(r"\d+\s+\d+\s+\d+(?:,\d+)?$", line):
                    flush()
                continue

        if buffer:
            buffer.append(line)
            if re.search(r"\d+\s+\d+\s+\d+(?:,\d+)?$", line):
                flush()

    flush()
    return items


def parse_site_products() -> list[dict]:
    doc = lxml_html.fromstring(cached_fetch(SITE_URL, "site-products.html"))
    products = []
    for product in doc.xpath('//*[@itemtype="http://schema.org/Product"]'):
        href = product.xpath('.//link[@itemprop="url"]/@href') or product.xpath(".//a/@href")
        detail_url = urljoin(SITE_BASE, href[0]) if href else SITE_URL
        detail_slug = slugify(detail_url.rstrip("/").split("/")[-1] or "product")
        detail = lxml_html.fromstring(cached_fetch(detail_url, f"detail-{detail_slug}.html"))
        name = " ".join(detail.xpath('//*[@itemprop="name"]//text()')[-1:]).strip()
        if not name:
            name = " ".join(product.xpath('.//*[@itemprop="name"]//text()')).strip()
        description = " ".join(" ".join(detail.xpath('//*[@itemprop="description"]//text()')).split())
        image = (detail.xpath('//img[@itemprop="image"]/@src') or product.xpath('.//img[@itemprop="image"]/@src') or [""])[0]
        products.append(
            {
                "name": name.strip(' "'),
                "description": description,
                "imageUrl": image,
                "url": detail_url,
                "brand": detect_brand(name),
                "normalized": normalize_text(name),
                "source": "famallnetwork",
            }
        )
    return products


def extract_window_data(source: bytes) -> dict:
    html = source.decode("utf-8", "replace")
    match = re.search(r"window\.data\s*=\s*(\{.*?\});\s*</script>", html, flags=re.S)
    if not match:
        return {}
    return json.loads(match.group(1))


def parse_taplink_products() -> list[dict]:
    main_data = extract_window_data(cached_fetch(TAPLINK_URL, "taplink-main.html"))
    collection_ids = []
    for field in main_data.get("fields", []):
        for item in field.get("items", []):
            options = item.get("options", {})
            if options.get("type") == "collection" and options.get("value"):
                collection_ids.append(str(format(int(options["value"]), "x")))
    collection_ids = list(dict.fromkeys(collection_ids))

    products = []
    seen_product_ids = set()
    for collection_id in collection_ids:
        collection_url = urljoin(TAPLINK_BASE, f"m/{collection_id}")
        data = extract_window_data(cached_fetch(collection_url, f"taplink-collection-{collection_id}.html"))
        collection_title = ""
        for collection in data.get("data", {}).get("collections", []):
            if str(collection.get("collection_id")) == collection_id:
                collection_title = collection.get("collection", "")
                break
        for product in data.get("data", {}).get("products", []):
            product_id = product.get("product_id")
            if product_id in seen_product_ids:
                continue
            seen_product_ids.add(product_id)
            name = str(product.get("title") or "").strip()
            picture = str(product.get("picture") or "").strip()
            image_url = urljoin(TAPLINK_IMAGE_BASE, picture) if picture else ""
            description = f"{name}. {collection_title}".strip()
            products.append(
                {
                    "name": name,
                    "description": description,
                    "imageUrl": image_url,
                    "url": collection_url,
                    "brand": detect_brand(name),
                    "normalized": normalize_text(name),
                    "source": "taplink",
                    "collection": collection_title,
                }
            )
    return products


def image_suffix(content_type: str) -> str:
    if "png" in content_type:
        return ".png"
    if "webp" in content_type:
        return ".webp"
    return ".jpg"


def extract_paragraph_images(paragraph, catalog_path: Path, index: int) -> list[str]:
    CATALOG_IMAGES_DIR.mkdir(parents=True, exist_ok=True)
    images = []
    for image_index, blip in enumerate(paragraph._element.xpath(".//a:blip"), start=1):
        relation_id = blip.get(qn("r:embed"))
        if not relation_id:
            continue
        image_part = paragraph.part.related_parts.get(relation_id)
        if not image_part:
            continue
        suffix = image_suffix(image_part.content_type)
        image_name = f"{slugify(catalog_path.stem)}-{index:03d}-{image_index}{suffix}"
        target = CATALOG_IMAGES_DIR / image_name
        if not target.exists():
            target.write_bytes(image_part.blob)
        images.append(str(target))
    return images


def parse_docx_catalog() -> list[dict]:
    catalog_paths = [path for path in CATALOG_DOCX_CANDIDATES if path.exists()]
    if not catalog_paths:
        return []

    products = []
    seen = set()
    for catalog_path in catalog_paths:
        doc = Document(catalog_path)
        pending_images: list[str] = []
        for paragraph_index, paragraph in enumerate(doc.paragraphs):
            paragraph_images = extract_paragraph_images(paragraph, catalog_path, paragraph_index)
            if paragraph_images:
                pending_images.extend(paragraph_images)

            lines = [line.strip() for line in paragraph.text.splitlines() if line.strip()]
            if not lines:
                continue

            title = lines[0]
            title_norm = normalize_text(title)
            if title.upper().startswith("FAMALL") or "каталог" in title_norm:
                continue
            if len(title) < 8:
                continue

            description_lines = [
                line
                for line in lines[1:]
                if not normalize_text(line).startswith("цена")
            ]
            description = " ".join(description_lines).strip()
            if not description and len(lines) == 1:
                description = title

            key = normalize_text(title)
            if key in seen:
                continue
            seen.add(key)
            image_path = pending_images.pop(0) if pending_images else ""
            products.append(
                {
                    "name": title,
                    "description": description,
                    "imageUrl": "",
                    "imagePath": image_path,
                    "url": str(catalog_path),
                    "brand": detect_brand(title),
                    "normalized": key,
                    "source": "docx_catalog",
                }
            )
    return products


def product_type_tokens(value: str) -> set[str]:
    tokens = set(normalize_text(value).split())
    type_groups = {
        "шампунь": {"шампунь", "шампуни"},
        "паста": {"паста", "зубная"},
        "салфетки": {"салфетки", "влажные"},
        "стирка": {"стирки", "стирка", "белья"},
        "трубы": {"труб", "трубы", "засоров"},
        "спрей": {"спрей"},
        "ополаскиватель": {"бальзам", "ополаскиватель"},
        "бомбочки": {"бомбочки", "ног"},
        "моющее": {"моющее", "посуды"},
        "чистящее": {"чистящее", "средство"},
    }
    found = set()
    for group, variants in type_groups.items():
        if tokens & variants:
            found.add(group)
    return found


def number_tokens(value: str) -> set[str]:
    return set(re.findall(r"\d+", normalize_text(value)))


def match_score(price_item: PriceItem, site_item: dict) -> tuple[float, str]:
    price_norm = normalize_text(price_item.name)
    site_norm = site_item["normalized"]
    if price_norm == site_norm:
        return 1.0, "exact_name"

    price_brand = normalize_text(price_item.brand)
    brand_match = price_item.brand != "FAMALL" and price_brand in site_norm
    price_types = product_type_tokens(price_item.name)
    site_types = product_type_tokens(site_item["name"])
    type_match = bool(price_types & site_types)
    price_numbers = number_tokens(f"{price_item.name} {price_item.volume}")
    site_numbers = number_tokens(site_item["name"])
    number_match = bool(price_numbers & site_numbers)
    if price_item.brand != "FAMALL" and not brand_match:
        return 0.0, "no_match"
    if price_types and site_types and not type_match:
        return 0.0, "no_match"
    if price_numbers and site_numbers and not number_match:
        return 0.0, "no_match"
    if price_item.brand == "FAMALL" and type_match and number_match:
        return 0.88, "type_volume"

    brand_score = 0.18 if brand_match else 0
    ratio = SequenceMatcher(None, price_norm, site_norm).ratio()
    price_tokens = set(price_norm.split())
    site_tokens = set(site_norm.split())
    meaningful = {token for token in price_tokens if len(token) > 4}
    token_score = len(meaningful & site_tokens) / max(len(meaningful), 1)
    score = max(ratio, token_score * 0.86) + brand_score

    if score >= 0.82 and (brand_match or type_match):
        return min(score, 0.99), "brand_keywords"
    if score >= 0.72 and brand_match and type_match:
        return score, "similar_name"
    return score, "no_match"


def download_image(url: str, slug: str) -> str:
    if not url:
        return ""
    suffix = ".png" if ".png" in url.lower().split("?")[0] else ".jpg"
    target = PRODUCT_IMAGES_DIR / f"{slug}{suffix}"
    if target.exists():
        return f"assets/images/products/{target.name}"
    try:
        data = fetch(url)
    except Exception:
        return ""
    target.write_bytes(data)
    return f"assets/images/products/{target.name}"


def use_source_image(source_item: dict, slug: str) -> str:
    image_path = source_item.get("imagePath")
    if image_path:
        source = Path(image_path)
        if source.exists():
            target = PRODUCT_IMAGES_DIR / f"{slug}{source.suffix.lower() or '.jpg'}"
            if not target.exists():
                shutil.copyfile(source, target)
            return f"assets/images/products/{target.name}"

    return download_image(source_item.get("imageUrl", ""), slug)


def create_premium_placeholder(item: PriceItem, slug: str) -> str:
    PLACEHOLDER_IMAGES_DIR.mkdir(parents=True, exist_ok=True)
    target = PLACEHOLDER_IMAGES_DIR / f"{slug}.svg"
    title_lines = wrap_text(item.name, line_limit=28, max_lines=4)
    title_markup = "\n".join(
        f'<text x="64" y="{230 + index * 34}" class="title">{xml_escape(line)}</text>'
        for index, line in enumerate(title_lines)
    )
    svg = f"""<svg xmlns="http://www.w3.org/2000/svg" width="900" height="900" viewBox="0 0 900 900" role="img" aria-label="{xml_escape(item.name)}">
  <defs>
    <linearGradient id="bg" x1="0" y1="0" x2="1" y2="1">
      <stop offset="0" stop-color="#ffffff"/>
      <stop offset="0.62" stop-color="#f7f8fa"/>
      <stop offset="1" stop-color="#fff0f2"/>
    </linearGradient>
    <filter id="shadow" x="-20%" y="-20%" width="140%" height="140%">
      <feDropShadow dx="0" dy="26" stdDeviation="30" flood-color="#202124" flood-opacity="0.14"/>
    </filter>
    <style>
      .brand {{ fill: #ff2942; font: 800 36px Arial, sans-serif; letter-spacing: 8px; }}
      .category {{ fill: #5f6368; font: 700 24px Arial, sans-serif; letter-spacing: 2px; }}
      .title {{ fill: #202124; font: 800 30px Arial, sans-serif; }}
      .note {{ fill: #5f6368; font: 600 22px Arial, sans-serif; }}
    </style>
  </defs>
  <rect width="900" height="900" rx="46" fill="url(#bg)"/>
  <g filter="url(#shadow)">
    <rect x="104" y="426" width="692" height="292" rx="44" fill="#ffffff"/>
    <rect x="134" y="456" width="632" height="232" rx="34" fill="#f7f8fa"/>
    <path d="M306 596c55-102 118-151 190-151 75 0 131 49 168 151" fill="none" stroke="#ff2942" stroke-width="18" stroke-linecap="round"/>
    <path d="M330 596h338" fill="none" stroke="#ff2942" stroke-width="18" stroke-linecap="round"/>
    <circle cx="390" cy="586" r="15" fill="#ff2942"/>
    <circle cx="586" cy="586" r="15" fill="#ff2942"/>
  </g>
  <text x="64" y="94" class="brand">FAMALL</text>
  <text x="64" y="156" class="category">{xml_escape(shorten(item.category, 38))}</text>
  {title_markup}
  <text x="64" y="805" class="note">Фото товара готовится</text>
  <text x="64" y="846" class="note">Цена из актуального PDF-прайса</text>
</svg>
"""
    target.write_text(svg, encoding="utf-8")
    return f"assets/images/products/placeholders/{target.name}"


def build_products() -> dict:
    PRODUCT_IMAGES_DIR.mkdir(parents=True, exist_ok=True)
    price_items = parse_price_pdf()
    docx_items = parse_docx_catalog()
    famallnetwork_items = parse_site_products()
    taplink_items = parse_taplink_products()
    site_items = famallnetwork_items + taplink_items
    content_items = docx_items + site_items
    used_image_indexes: set[int] = set()
    used_content_indexes: set[int] = set()
    manual_review = []
    price_without_content = []
    products = []

    for item in price_items:
        best_image_index = -1
        best_image_score = 0.0
        best_image_method = "no_match"
        for index, source_item in enumerate(content_items):
            if index in used_image_indexes or not (source_item.get("imageUrl") or source_item.get("imagePath")):
                continue
            score, method = match_score(item, source_item)
            if score > best_image_score:
                best_image_index = index
                best_image_score = score
                best_image_method = method

        best_description_index = -1
        best_description_score = 0.0
        best_description_method = "no_match"
        for index, source_item in enumerate(content_items):
            score, method = match_score(item, source_item)
            if score > best_description_score:
                best_description_index = index
                best_description_score = score
                best_description_method = method

        matched_image = (
            content_items[best_image_index]
            if best_image_index >= 0 and best_image_method in {"exact_name", "brand_keywords", "type_volume"}
            else None
        )
        matched_description = (
            content_items[best_description_index]
            if best_description_index >= 0 and best_description_method in {"exact_name", "brand_keywords", "type_volume"}
            else None
        )
        slug = slugify(f"{item.index}-{item.name}")
        image = ""
        image_type = "source"
        description = ""
        status = "published"
        match_info = {
            "method": "none",
            "score": round(max(best_image_score, best_description_score), 3),
            "sourceName": None,
        }

        if matched_image:
            used_image_indexes.add(best_image_index)
            used_content_indexes.add(best_image_index)
            image = use_source_image(matched_image, slug)
            description = matched_image["description"]
            match_info = {
                "method": best_image_method,
                "score": round(best_image_score, 3),
                "sourceName": matched_image["name"],
                "source": matched_image.get("source"),
            }
        elif best_image_index >= 0 and best_image_method == "similar_name":
            manual_review.append(
                {
                    "priceName": item.name,
                    "sourceName": content_items[best_image_index]["name"],
                    "source": content_items[best_image_index].get("source"),
                    "score": round(best_image_score, 3),
                    "reason": "Похожее название найдено, но фото не опубликовано до ручной проверки.",
                }
            )

        if not image:
            image = create_premium_placeholder(item, slug)
            image_type = "premium_placeholder"
            status = "need_image"

        if matched_description and matched_description.get("description"):
            used_content_indexes.add(best_description_index)
            description = matched_description["description"]

        if not description:
            description = f"{item.name}. Описание будет дополнено после загрузки полного каталога FAMALL."

        if not matched_image and not matched_description:
            price_without_content.append(
                {
                    "id": f"famall-{item.index:03d}",
                    "name": item.name,
                    "category": item.category,
                    "reason": "Товар есть в PDF-прайсе, но не найден в каталоге, famallnetwork62.orgs.biz и Taplink. Карточка опубликована с placeholder.",
                }
            )

        products.append(
            {
                "id": f"famall-{item.index:03d}",
                "sku": f"FAMALL-{item.index:03d}",
                "slug": slug,
                "name": item.name,
                "brand": item.brand,
                "category": item.category,
                "volume": item.volume,
                "description": description,
                "retailPrice": item.retail_price,
                "partnerPrice": item.partner_price,
                "pv": item.pv,
                "image": image,
                "imageType": image_type,
                "gallery": [
                    {
                        "type": image_type,
                        "src": image,
                        "alt": item.name,
                    }
                ],
                "characteristics": [
                    {"label": "Артикул", "value": f"FAMALL-{item.index:03d}"},
                    {"label": "Бренд", "value": item.brand},
                    {"label": "Категория", "value": item.category},
                    {"label": "Объём", "value": item.volume or "уточняется"},
                    {"label": "PV", "value": str(item.pv).replace(".", ",")},
                    {"label": "Наличие", "value": "В наличии"},
                ],
                "discountPercent": round((1 - item.partner_price / item.retail_price) * 100) if item.retail_price else 0,
                "collections": [],
                "badges": [],
                "stock": "in_stock",
                "status": status,
                "match": match_info,
            }
        )

    for product in products[-24:]:
        product["collections"].append("new")
        product["badges"].append("Новинка")

    for product in sorted(products, key=lambda entry: entry["discountPercent"], reverse=True)[:24]:
        if "sale" not in product["collections"]:
            product["collections"].append("sale")
        if "Акция" not in product["badges"]:
            product["badges"].append("Акция")

    real_image_products = [product for product in products if product["imageType"] == "source"]
    hit_candidates = [
        product
        for product in real_image_products
        if normalize_text(product["category"]) in {
            normalize_text("КУХНЯ"),
            normalize_text("ДЛЯ ПОЛОСТИ РТА"),
            normalize_text("УХОД ЗА ТЕЛОМ"),
            normalize_text("УХОДОВЫЕ ШАМПУНИ"),
            normalize_text("ДЕКОРАТИВНАЯ КОСМЕТИКА"),
        }
    ]
    for product in (hit_candidates or real_image_products)[:24]:
        if "hits" not in product["collections"]:
            product["collections"].append("hits")
        if "Хит" not in product["badges"]:
            product["badges"].append("Хит")

    unmatched = [content_items[index] for index in range(len(content_items)) if index not in used_content_indexes]
    for item in unmatched:
        item.pop("normalized", None)

    missing_images = [
        {
            "id": product["id"],
            "sku": product["sku"],
            "name": product["name"],
            "brand": product["brand"],
            "category": product["category"],
            "volume": product["volume"],
            "retailPrice": product["retailPrice"],
            "partnerPrice": product["partnerPrice"],
            "pv": product["pv"],
            "placeholderImage": product["image"],
            "status": product["status"],
            "sourcesTried": [
                "DOCX-каталог без цен",
                SITE_URL,
                TAPLINK_URL,
            ],
            "result": "Создана премиальная карточка-заглушка",
            "nextAction": "Нужна ручная загрузка реального фото или AI-генерация изображения товара.",
        }
        for product in products
        if product["imageType"] != "source"
    ]

    report = {
        "priceSource": str(PRICE_PDF),
        "pricePolicy": "Цены, PV и публикация товаров берутся только из PDF-прайса.",
        "catalogPolicy": "DOCX-каталог и сайты используются только для названий, описаний, категорий и изображений. Отсутствие цены в каталоге не является ошибкой.",
        "docxCatalogSourcesChecked": [str(path) for path in CATALOG_DOCX_CANDIDATES],
        "docxCatalogSourcesUsed": [str(path) for path in CATALOG_DOCX_CANDIDATES if path.exists()],
        "docxCatalogProductsFound": len(docx_items),
        "siteSources": [SITE_URL, TAPLINK_URL],
        "priceProductsFound": len(price_items),
        "famallnetworkProductsFound": len(famallnetwork_items),
        "taplinkProductsFound": len(taplink_items),
        "siteProductsFound": len(site_items),
        "productsPublished": len(products),
        "productsMatchedWithSources": len(price_items) - len(price_without_content),
        "productsMatchedWithImage": sum(1 for item in products if item["imageType"] == "source"),
        "productsWithDescriptions": sum(
            1
            for item in products
            if "Описание будет дополнено" not in item["description"]
        ),
        "productsWithoutImage": len(missing_images),
        "premiumPlaceholdersCreated": len(missing_images),
        "newProductsPageCount": sum(1 for item in products if "new" in item["collections"]),
        "hitProductsPageCount": sum(1 for item in products if "hits" in item["collections"]),
        "saleProductsPageCount": sum(1 for item in products if "sale" in item["collections"]),
        "catalogProductsNotPublished": len(unmatched),
        "siteProductsNotPublished": sum(1 for item in unmatched if item.get("source") != "docx_catalog"),
        "docxProductsNotPublished": sum(1 for item in unmatched if item.get("source") == "docx_catalog"),
        "priceProductsWithoutContentMatch": len(price_without_content),
        "priceProductsWithoutWebMatch": len(price_without_content),
        "manualReviewCount": len(manual_review),
        "manualReview": manual_review,
        "priceProductsWithoutContentMatchList": price_without_content,
        "priceProductsWithoutWebMatchList": price_without_content,
    }
    review = {
        "similarMatches": manual_review,
        "priceProductsWithoutContentMatch": price_without_content,
        "priceProductsWithoutWebMatch": price_without_content,
    }
    return {"products": products, "unmatched": unmatched, "report": report, "review": review, "missingImages": missing_images}


def main() -> None:
    result = build_products()
    PRODUCTS_JSON.write_text(json.dumps(result["products"], ensure_ascii=False, indent=2), encoding="utf-8")
    categories = []
    seen_categories = set()
    for product in result["products"]:
        category = product["category"]
        if category in seen_categories:
            continue
        seen_categories.add(category)
        categories.append({"id": slugify(category), "name": category, "slug": slugify(category)})
    CATEGORIES_JSON.write_text(json.dumps(categories, ensure_ascii=False, indent=2), encoding="utf-8")
    UNMATCHED_JSON.write_text(json.dumps(result["unmatched"], ensure_ascii=False, indent=2), encoding="utf-8")
    REVIEW_JSON.write_text(json.dumps(result["review"], ensure_ascii=False, indent=2), encoding="utf-8")
    REPORT_JSON.write_text(json.dumps(result["report"], ensure_ascii=False, indent=2), encoding="utf-8")
    MISSING_IMAGES_JSON.write_text(json.dumps(result["missingImages"], ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result["report"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
